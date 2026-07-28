"""
Security and functionality verification tests for backend security controls.
"""

import sys
from pathlib import Path
sys.path.append(str(Path(__file__).parent))

from fastapi.testclient import TestClient
from main import app
import docx

client = TestClient(app)


def test_security_and_functionality():
    print("=" * 60)
    print("TESTING BACKEND SECURITY CONTROLS & FUNCTIONALITY")
    print("=" * 60)

    # 1. Health check
    response = client.get("/health")
    assert response.status_code == 200
    print("[SUCCESS] GET /health PASSED")

    # 2. Valid DOCX analysis and JD match
    test_docx_path = Path(__file__).parent / "test_resumes" / "sample_api_test.docx"
    doc = docx.Document()
    doc.add_heading("Jane Security Test", 0)
    doc.add_paragraph("Email: jane.security@tech.io | Phone: +1 555-000-1111")
    doc.add_heading("Skills", 1)
    doc.add_paragraph("Python, FastAPI, React, PostgreSQL, Docker, AWS, Git")
    doc.add_heading("Work Experience", 1)
    doc.add_paragraph("Senior Security Engineer - SafeTech | 2021 - Present")
    doc.add_paragraph("• Implemented security validation controls and REST APIs.")
    doc.save(str(test_docx_path))

    with open(test_docx_path, "rb") as f:
        response = client.post("/analyze", files={"file": ("../../dangerous_filename<script>.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert response.status_code == 200, f"Valid DOCX failed: {response.text}"
    rec = response.json()
    assert rec["file_name"] == "dangerous_filename_script_.docx", f"Filename sanitization failed: {rec['file_name']}"
    print(f"[SUCCESS] Path traversal & symbol sanitization PASSED ({rec['file_name']})")
    print(f"[SUCCESS] POST /analyze PASSED (ATS Score: {rec['ats_score']})")

    # 3. Test JD Match
    analysis_id = rec["id"]
    jd_payload = {
        "analysis_id": analysis_id,
        "job_description": "Seeking Senior Security Engineer proficient in Python, FastAPI, and PostgreSQL."
    }
    response = client.post("/match-jd", json=jd_payload)
    assert response.status_code == 200
    match_res = response.json()
    assert match_res["match_result"]["match_percentage"] > 50
    print(f"[SUCCESS] POST /match-jd PASSED (Match score: {match_res['match_result']['match_percentage']}%)")

    # 4. Security Test: Invalid Magic Bytes (Fake PDF)
    fake_pdf_content = b"This is plain text pretending to be a pdf file."
    response = client.post("/upload", files={"file": ("fake_resume.pdf", fake_pdf_content, "application/pdf")})
    assert response.status_code == 400, "Fake PDF should have been rejected with 400"
    assert "MIME type validation failed" in response.json()["detail"]
    print("[SUCCESS] Fake PDF magic-byte validation PASSED (Rejected with HTTP 400)")

    # 5. Security Test: Oversized File (>10MB)
    large_fake_content = b"%PDF-" + b"0" * (11 * 1024 * 1024)
    response = client.post("/upload", files={"file": ("large_resume.pdf", large_fake_content, "application/pdf")})
    assert response.status_code == 400, "Oversized file should have been rejected with 400"
    assert "exceeds maximum allowed size limit of 10MB" in response.json()["detail"]
    print("[SUCCESS] Oversized file limit PASSED (Rejected with HTTP 400)")

    print("\nALL BACKEND SECURITY & FUNCTIONALITY TESTS PASSED SUCCESSFULLY!")


if __name__ == "__main__":
    test_security_and_functionality()
