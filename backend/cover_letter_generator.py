"""
Cover Letter Generation Engine.
Generates authentic, high-impact, ATS-aligned cover letters
using Groq LLM with intelligent offline NLP rule-based fallback.
Supports 1-click DOCX export.
"""

import os
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional
import docx
from dotenv import load_dotenv

# Ensure environment variables are loaded
_backend_env = Path(__file__).parent / ".env"
_root_env = Path(__file__).parent.parent / ".env"
load_dotenv(_backend_env)
load_dotenv(_root_env)

from resume_parser import SKILL_TAXONOMY


class CoverLetterGenerator:
    """Production Cover Letter Generator & Formatter."""

    def __init__(self, api_key: Optional[str] = None):
        self.groq_api_key = api_key or os.getenv("GROQ_API_KEY")
        self.openai_api_key = os.getenv("OPENAI_API_KEY")

    def generate(
        self,
        parsed_resume: Dict[str, Any],
        job_description: str,
        company_name: Optional[str] = None,
        job_title: Optional[str] = None,
        tone: str = "professional"
    ) -> Dict[str, Any]:
        """
        Generates tailored cover letter combining candidate verified background with target role.
        """
        candidate_name = parsed_resume.get("name") or "Candidate"
        candidate_email = parsed_resume.get("email") or ""
        candidate_phone = parsed_resume.get("phone") or ""
        candidate_location = parsed_resume.get("location") or ""
        skills = parsed_resume.get("skills") or []
        experiences = parsed_resume.get("experience") or []
        projects = parsed_resume.get("projects") or []

        resolved_company = company_name.strip() if company_name and company_name.strip() else "the hiring organization"
        resolved_title = job_title.strip() if job_title and job_title.strip() else self._infer_job_title(job_description)

        # 1. Identify matching & relevant skills
        jd_skills = self._extract_jd_skills(job_description)
        candidate_skills_lower = set([s.lower() for s in skills])
        matching_skills = [s.title() for s in jd_skills.intersection(candidate_skills_lower)]

        # 2. Try LLM Generation if API key is provided
        if self.groq_api_key or self.openai_api_key:
            llm_result = self._try_llm_cover_letter(
                candidate_name=candidate_name,
                skills=skills,
                experiences=experiences,
                projects=projects,
                job_description=job_description,
                company_name=resolved_company,
                job_title=resolved_title,
                tone=tone
            )
            if llm_result:
                return {
                    "cover_letter_text": llm_result,
                    "company_name": resolved_company,
                    "job_title": resolved_title,
                    "tone": tone,
                    "matching_skills_highlighted": matching_skills[:6],
                    "source": "llm"
                }

        # 3. Smart Offline Rule-Based Generator
        offline_result = self._generate_rule_based_cover_letter(
            candidate_name=candidate_name,
            candidate_email=candidate_email,
            candidate_phone=candidate_phone,
            candidate_location=candidate_location,
            skills=skills,
            experiences=experiences,
            projects=projects,
            matching_skills=matching_skills,
            company_name=resolved_company,
            job_title=resolved_title,
            tone=tone
        )

        return {
            "cover_letter_text": offline_result,
            "company_name": resolved_company,
            "job_title": resolved_title,
            "tone": tone,
            "matching_skills_highlighted": matching_skills[:6],
            "source": "smart_nlp_engine"
        }

    def _infer_job_title(self, jd_text: str) -> str:
        """Heuristic to extract or infer job title from JD text."""
        patterns = [
            r'(?:seeking|hiring|looking for)(?:\s+an?|\s+our)?\s+([A-Z][a-zA-Z\s]{2,30}(?:Engineer|Developer|Manager|Architect|Analyst|Specialist|Lead|Designer))',
            r'\b((?:Senior|Lead|Principal|Junior|Staff)?\s*(?:Full Stack|Frontend|Backend|Software|DevOps|Data|Cloud|AI|ML)\s*(?:Engineer|Developer|Architect))\b'
        ]
        for pat in patterns:
            m = re.search(pat, jd_text, re.IGNORECASE)
            if m:
                return m.group(1).strip()
        return "Software Engineer"

    def _extract_jd_skills(self, jd_text: str) -> set:
        found = set()
        jd_lower = jd_text.lower()
        for skill in SKILL_TAXONOMY:
            if re.search(r'\b' + re.escape(skill) + r'\b', jd_lower):
                found.add(skill.lower())
        return found

    def _try_llm_cover_letter(
        self,
        candidate_name: str,
        skills: List[str],
        experiences: List[Dict[str, Any]],
        projects: List[Dict[str, Any]],
        job_description: str,
        company_name: str,
        job_title: str,
        tone: str
    ) -> Optional[str]:
        """Call Groq or OpenAI API for tailored cover letter."""
        system_prompt = f"""You are an elite, executive career strategist and technical cover letter writer.
Your job is to author a compelling, modern, and highly persuasive cover letter for a candidate applying for the role of {job_title} at {company_name}.

CRITICAL RULES:
1. TONE: {tone.capitalize()}, authentic, professional, and confident. Avoid clichés (e.g., "I am writing to express my enthusiasm", "thrilled to apply").
2. AUTHENTICITY GUARDRAIL: Only reference technologies, companies, and achievements genuinely present in the candidate's verified profile. Do not invent false metrics or degrees.
3. STRUCTURE:
   - Header with Salutation: Dear Hiring Team at {company_name},
   - Paragraph 1: Direct, compelling hook connecting the candidate's core expertise to {company_name}'s mission/role.
   - Paragraph 2: Specific technical accomplishments from recent experience, highlighting architecture, problem-solving, and deliverables.
   - Paragraph 3: Alignment with target job description requirements and technical competencies.
   - Paragraph 4: Enthusiastic closing highlighting immediate contribution and call to interview.
   - Sign-off: Sincerely, followed by {candidate_name}.
4. OUTPUT FORMAT: Return ONLY the clean, formatted letter body with line breaks. Do not wrap in markdown quotes or preamble."""

        exp_summary = []
        for exp in experiences[:2]:
            t = exp.get("title", "Role")
            c = exp.get("company", "Company")
            desc = exp.get("description", "")[:200]
            exp_summary.append(f"- {t} at {c}: {desc}")

        user_prompt = f"""Candidate Name: {candidate_name}
Target Role: {job_title}
Target Company: {company_name}
Candidate Key Skills: {', '.join(skills[:12])}
Recent Experience:
{chr(10).join(exp_summary)}
Job Description Overview:
{job_description[:600]}

Generate the complete tailored cover letter now:"""

        # Try Groq API
        if self.groq_api_key:
            try:
                from groq import Groq
                client = Groq(api_key=self.groq_api_key)
                groq_model = os.getenv("GROQ_MODEL", "openai/gpt-oss-120b")
                completion = client.chat.completions.create(
                    model=groq_model,
                    messages=[
                        {"role": "system", "content": system_prompt},
                        {"role": "user", "content": user_prompt}
                    ],
                    temperature=0.3
                )
                text = completion.choices[0].message.content.strip()
                if len(text) > 200:
                    return text
            except Exception as e:
                print(f"[COVER LETTER LLM] Groq API call failed: {e}")

        # Try OpenAI API
        if self.openai_api_key:
            try:
                import openai
                client = openai.OpenAI(api_key=self.openai_api_key)
                completion = client.chat.completions.create(
                    model="gpt-4o-mini",
                    messages=[
                        {"role": "system", "content": system_prompt},
                        {"role": "user", "content": user_prompt}
                    ],
                    temperature=0.3
                )
                text = completion.choices[0].message.content.strip()
                if len(text) > 200:
                    return text
            except Exception as e:
                print(f"[COVER LETTER LLM] OpenAI API call failed: {e}")

        return None

    def _generate_rule_based_cover_letter(
        self,
        candidate_name: str,
        candidate_email: str,
        candidate_phone: str,
        candidate_location: str,
        skills: List[str],
        experiences: List[Dict[str, Any]],
        projects: List[Dict[str, Any]],
        matching_skills: List[str],
        company_name: str,
        job_title: str,
        tone: str
    ) -> str:
        """Smart deterministic cover letter composer."""
        today_str = datetime.now().strftime("%B %d, %Y")

        top_skills_str = ", ".join(matching_skills[:5]) if matching_skills else ", ".join(skills[:5])

        # Extract recent role highlight
        recent_role = experiences[0] if experiences else None
        role_title = recent_role.get("title", "Software Engineer") if recent_role else "Software Engineer"
        role_company = recent_role.get("company", "") if recent_role else ""
        company_phrase = f" at {role_company}" if role_company else ""

        # Extract bullet highlight
        desc_highlight = ""
        if recent_role and recent_role.get("description"):
            lines = [l.strip().lstrip('•-* ') for l in recent_role["description"].split('\n') if len(l.strip()) > 15]
            if lines:
                desc_highlight = lines[0]

        # Extract project highlight
        project_highlight = ""
        if projects and projects[0].get("title"):
            p = projects[0]
            project_highlight = f"In addition, through my work on {p.get('title')}, I designed and deployed scalable technical deliverables aligned with production best practices."

        letter_parts = [
            f"Dear Hiring Team at {company_name},\n",
            f"I am writing to submit my application for the {job_title} position at {company_name}. With proven engineering experience as a {role_title}{company_phrase} and a deep technical background in {top_skills_str}, I have consistently architected, deployed, and scaled high-performance software systems that solve critical business challenges.\n",
            f"In my recent role as {role_title}{company_phrase}, I led core engineering initiatives emphasizing code quality, architectural resilience, and scalable delivery. {desc_highlight} By leveraging {top_skills_str}, I specialize in bridging technical implementation with end-user performance, ensuring systems remain reliable, maintainable, and aligned with modern industry benchmarks.\n",
            f"{project_highlight} What excites me most about {company_name} is the opportunity to bring my hands-on problem-solving skills and domain expertise in {matching_skills[0] if matching_skills else 'software engineering'} to your engineering organization, collaborating with cross-functional partners to drive meaningful product impact from Day 1.\n",
            f"I welcome the opportunity to discuss how my technical competencies, execution velocity, and engineering mindset can contribute directly to {company_name}'s goals. Thank you for your time and consideration.\n",
            f"Sincerely,\n{candidate_name}"
        ]

        return "\n".join(letter_parts)

    def export_docx(
        self,
        candidate_name: str = "Candidate",
        candidate_email: str = "",
        candidate_phone: str = "",
        candidate_location: str = "",
        company_name: str = "Hiring Organization",
        job_title: str = "Target Role",
        cover_letter_text: str = ""
    ) -> docx.Document:
        """Exports formatted DOCX document of the cover letter."""
        doc = docx.Document()

        # Header: Candidate Info
        doc.add_heading(candidate_name, level=0)

        contact_parts = []
        if candidate_email: contact_parts.append(candidate_email)
        if candidate_phone: contact_parts.append(candidate_phone)
        if candidate_location: contact_parts.append(candidate_location)
        if contact_parts:
            doc.add_paragraph(" | ".join(contact_parts))

        doc.add_paragraph("")  # Spacing

        # Date & Recipient
        today_str = datetime.now().strftime("%B %d, %Y")
        doc.add_paragraph(today_str)
        doc.add_paragraph(f"Hiring Team\n{company_name}\nRe: Application for {job_title}")
        doc.add_paragraph("")  # Spacing

        # Letter Body
        paragraphs = [p.strip() for p in cover_letter_text.split("\n\n") if p.strip()]
        for p in paragraphs:
            doc.add_paragraph(p)

        return doc
