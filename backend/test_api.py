"""
API integration test for FastAPI backend endpoints.
"""

import sys
from pathlib import Path
sys.path.append(str(Path(__file__).parent))

from fastapi.testclient import TestClient
from main import app
import docx

client = TestClient(app)


def test_api_endpoints():
    print("=" * 60)
    print("TESTING FASTAPI BACKEND API ENDPOINTS")
    print("=" * 60)

    # 1. Health check
    response = client.get("/health")
    assert response.status_code == 200, "Health check failed"
    print("[SUCCESS] GET /health PASSED")

    # Generate sample docx file
    test_docx_path = Path(__file__).parent / "test_resumes" / "sample_api_test.docx"
    doc = docx.Document()
    doc.add_heading("Jane Doe", 0)
    doc.add_paragraph("Email: jane.doe@tech.io | Phone: +1 555-999-8888")
    doc.add_heading("Skills", 1)
    doc.add_paragraph("Python, FastAPI, React, PostgreSQL, Docker, AWS, Git")
    doc.add_heading("Work Experience", 1)
    doc.add_paragraph("Senior Backend Engineer - Tech Solutions | 2021 - Present")
    doc.add_paragraph("• Architected FastAPI microservices handling millions of API calls daily.")
    doc.add_heading("Education", 1)
    doc.add_paragraph("Bachelor of Science in Software Engineering - MIT | 2017 - 2021")
    doc.save(str(test_docx_path))

    # 2. Upload Endpoint
    with open(test_docx_path, "rb") as f:
        response = client.post("/upload", files={"file": ("sample_api_test.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert response.status_code == 200, f"Upload failed: {response.text}"
    upload_res = response.json()
    assert upload_res["success"] is True
    print("[SUCCESS] POST /upload PASSED")

    # Register a test user for authenticated endpoints
    import time
    ts = int(time.time())
    reg_res = client.post("/auth/register", json={
        "email": f"api_test_{ts}@example.com",
        "password": "SecurePassword123!",
        "full_name": "API Test User"
    })
    assert reg_res.status_code == 200, f"Registration failed: {reg_res.text}"
    token = reg_res.json()["access_token"]
    auth_headers = {"Authorization": f"Bearer {token}"}
    print("[SUCCESS] POST /auth/register PASSED")

    # 3. Analyze Endpoint (Authenticated)
    with open(test_docx_path, "rb") as f:
        response = client.post(
            "/analyze",
            files={"file": ("sample_api_test.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            headers=auth_headers
        )
    assert response.status_code == 200, f"Analyze failed: {response.text}"
    analyze_res = response.json()
    analysis_id = analyze_res["id"]
    print(f"[SUCCESS] POST /analyze PASSED (Created Analysis ID: {analysis_id}, ATS Score: {analyze_res['ats_score']})")

    # 4. Match JD Endpoint
    jd_payload = {
        "analysis_id": analysis_id,
        "job_description": "We are seeking a Senior Backend Engineer proficient in Python, FastAPI, PostgreSQL, Docker, and AWS."
    }
    response = client.post("/match-jd", json=jd_payload, headers=auth_headers)
    assert response.status_code == 200, f"Match JD failed: {response.text}"
    match_res = response.json()
    print(f"[SUCCESS] POST /match-jd PASSED (Match score: {match_res['match_result']['match_percentage']}%)")

    # 5. Get History Endpoint
    response = client.get("/history", headers=auth_headers)
    assert response.status_code == 200, f"History failed: {response.text}"
    history = response.json()
    assert len(history) > 0
    print(f"[SUCCESS] GET /history PASSED ({len(history)} records found)")

    # 6. Get Analysis by ID
    response = client.get(f"/analysis/{analysis_id}", headers=auth_headers)
    assert response.status_code == 200, f"Get analysis failed: {response.text}"
    rec = response.json()
    assert rec["id"] == analysis_id
    print(f"[SUCCESS] GET /analysis/{analysis_id} PASSED")

    # 7. Rewrite Resume Endpoint
    rw_payload = {
        "analysis_id": analysis_id,
        "job_description": "We are seeking a Senior Backend Engineer proficient in Python, FastAPI, and PostgreSQL."
    }
    response = client.post("/rewrite-resume", json=rw_payload, headers=auth_headers)
    assert response.status_code == 200, f"Rewrite resume failed: {response.text}"
    rw_res = response.json()
    assert "rewrites" in rw_res and len(rw_res["rewrites"]) > 0
    print(f"[SUCCESS] POST /rewrite-resume PASSED ({len(rw_res['rewrites'])} rewrites generated)")

    # 8. Export Tailored Resume (.docx)
    exp_payload = {
        "analysis_id": analysis_id,
        "accepted_rewrites": rw_res["rewrites"]
    }
    response = client.post("/export-tailored-resume", json=exp_payload, headers=auth_headers)
    assert response.status_code == 200, f"Export tailored resume failed: {response.text}"
    assert len(response.content) > 1000
    print(f"[SUCCESS] POST /export-tailored-resume PASSED ({len(response.content)} bytes)")

    # 9. Generate Cover Letter Endpoint
    cl_payload = {
        "analysis_id": analysis_id,
        "job_description": "We are seeking a Senior Backend Engineer proficient in Python, FastAPI, and PostgreSQL at Stripe.",
        "company_name": "Stripe",
        "job_title": "Senior Backend Engineer"
    }
    response = client.post("/generate-cover-letter", json=cl_payload, headers=auth_headers)
    assert response.status_code == 200, f"Generate cover letter failed: {response.text}"
    cl_res = response.json()
    assert cl_res["success"] is True and len(cl_res["cover_letter"]["cover_letter_text"]) > 200
    print(f"[SUCCESS] POST /generate-cover-letter PASSED ({len(cl_res['cover_letter']['cover_letter_text'])} chars)")

    # 10. Export Cover Letter (.docx)
    exp_cl_payload = {
        "analysis_id": analysis_id,
        "company_name": "Stripe",
        "job_title": "Senior Backend Engineer",
        "cover_letter_text": cl_res["cover_letter"]["cover_letter_text"]
    }
    response = client.post("/export-cover-letter", json=exp_cl_payload, headers=auth_headers)
    assert response.status_code == 200, f"Export cover letter failed: {response.text}"
    assert len(response.content) > 1000
    print(f"[SUCCESS] POST /export-cover-letter PASSED ({len(response.content)} bytes)")

    # 11. Generate Interview Prep Endpoint
    ip_payload = {
        "analysis_id": analysis_id,
        "job_description": "We are seeking a Senior Backend Engineer proficient in Python, FastAPI, React, and Kubernetes at Stripe."
    }
    response = client.post("/interview-prep", json=ip_payload, headers=auth_headers)
    assert response.status_code == 200, f"Interview prep failed: {response.text}"
    ip_res = response.json()
    assert ip_res["success"] is True and len(ip_res["interview_prep"]["questions"]) > 0
    print(f"[SUCCESS] POST /interview-prep PASSED ({len(ip_res['interview_prep']['questions'])} questions generated)")

    # 12. Export Application Bundle (ZIP) Endpoint
    bundle_payload = {
        "analysis_id": analysis_id,
        "company_name": "Stripe",
        "job_title": "Senior Backend Engineer",
        "cover_letter_text": cl_res["cover_letter"]["cover_letter_text"]
    }
    response = client.post("/export-application-bundle", json=bundle_payload, headers=auth_headers)
    assert response.status_code == 200, f"Export application bundle failed: {response.text}"
    assert response.headers.get("content-type") == "application/zip"
    assert len(response.content) > 2000

    import zipfile
    from io import BytesIO
    with zipfile.ZipFile(BytesIO(response.content)) as z:
        names = z.namelist()
        assert any(n.endswith(".docx") and "Resume" in n for n in names)
        assert any(n.endswith(".docx") and "Cover_Letter" in n for n in names)
        assert any(n.endswith(".md") and "Interview_Prep" in n for n in names)
    print(f"[SUCCESS] POST /export-application-bundle PASSED ({len(response.content)} bytes, {len(names)} files in ZIP: {names})")


    print("\nALL FASTAPI ENDPOINT TESTS PASSED SUCCESSFULLY!")


if __name__ == "__main__":
    test_api_endpoints()
