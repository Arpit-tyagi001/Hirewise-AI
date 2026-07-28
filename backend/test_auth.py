"""
Authentication and Data Access Isolation test suite.
Verifies User A vs User B data isolation, JWT authentication, and anonymous demo mode.
"""

import time
import sys
from pathlib import Path
sys.path.append(str(Path(__file__).parent))

from fastapi.testclient import TestClient
from main import app
import docx

client = TestClient(app)


def main():
    print("=" * 60)
    print("TESTING JWT AUTHENTICATION & DATA ACCESS ISOLATION")
    print("=" * 60)

    # Generate test document
    test_docx_path = Path(__file__).parent / "test_resumes" / "auth_test.docx"
    test_docx_path.parent.mkdir(exist_ok=True)
    doc = docx.Document()
    doc.add_heading("Auth Test Candidate", 0)
    doc.add_paragraph("Email: test.candidate@tech.io | Phone: +1 555-123-4567")
    doc.add_heading("Skills", 1)
    doc.add_paragraph("Python, FastAPI, React, PostgreSQL, Docker, AWS")
    doc.add_heading("Experience", 1)
    doc.add_paragraph("Software Engineer - Tech Enterprise | 2021 - Present")
    doc.add_paragraph("• Built authenticated microservices and REST APIs.")
    doc.save(str(test_docx_path))

    timestamp = int(time.time())
    user_a_email = f"usera_{timestamp}@tech.io"
    user_a_pass = "Password123!"

    user_b_email = f"userb_{timestamp}@tech.io"
    user_b_pass = "Password456!"

    # 1. Register User A
    res_a = client.post("/auth/register", json={"email": user_a_email, "password": user_a_pass, "full_name": "Alice UserA"})
    assert res_a.status_code == 200, f"User A registration failed: {res_a.text}"
    token_a = res_a.json()["access_token"]
    user_a_id = res_a.json()["user"]["id"]
    headers_a = {"Authorization": f"Bearer {token_a}"}
    print(f"[SUCCESS] User A Registered (ID: {user_a_id}, Email: {user_a_email})")

    # 2. Register User B
    res_b = client.post("/auth/register", json={"email": user_b_email, "password": user_b_pass, "full_name": "Bob UserB"})
    assert res_b.status_code == 200, f"User B registration failed: {res_b.text}"
    token_b = res_b.json()["access_token"]
    user_b_id = res_b.json()["user"]["id"]
    headers_b = {"Authorization": f"Bearer {token_b}"}
    print(f"[SUCCESS] User B Registered (ID: {user_b_id}, Email: {user_b_email})")

    # 3. Test Login
    login_res = client.post("/auth/login", json={"email": user_a_email, "password": user_a_pass})
    assert login_res.status_code == 200
    assert "access_token" in login_res.json()
    print("[SUCCESS] User A Login PASSED")

    # 4. Anonymous Demo Mode Upload/Analyze (No Authorization Header)
    with open(test_docx_path, "rb") as f:
        anon_res = client.post("/analyze", files={"file": ("anon_resume.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert anon_res.status_code == 200, f"Anon analyze failed: {anon_res.text}"
    assert anon_res.json()["user_id"] is None
    print("[SUCCESS] Anonymous Demo Mode /analyze PASSED (user_id=None)")

    # 5. User A Upload/Analyze
    with open(test_docx_path, "rb") as f:
        user_a_rec_res = client.post("/analyze", headers=headers_a, files={"file": ("alice_resume.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert user_a_rec_res.status_code == 200, f"User A analyze failed: {user_a_rec_res.text}"
    record_a_id = user_a_rec_res.json()["id"]
    assert user_a_rec_res.json()["user_id"] == user_a_id
    print(f"[SUCCESS] User A /analyze PASSED (Created record ID #{record_a_id} linked to User A)")

    # 6. User B Upload/Analyze
    with open(test_docx_path, "rb") as f:
        user_b_rec_res = client.post("/analyze", headers=headers_b, files={"file": ("bob_resume.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert user_b_rec_res.status_code == 200, f"User B analyze failed: {user_b_rec_res.text}"
    record_b_id = user_b_rec_res.json()["id"]
    assert user_b_rec_res.json()["user_id"] == user_b_id
    print(f"[SUCCESS] User B /analyze PASSED (Created record ID #{record_b_id} linked to User B)")

    # 7. Test Data Isolation on GET /history
    history_a = client.get("/history", headers=headers_a).json()
    history_b = client.get("/history", headers=headers_b).json()
    
    ids_a = [r["id"] for r in history_a]
    ids_b = [r["id"] for r in history_b]

    assert record_a_id in ids_a, "User A should see Record A"
    assert record_b_id not in ids_a, "User A MUST NOT see Record B"
    assert record_b_id in ids_b, "User B should see Record B"
    assert record_a_id not in ids_b, "User B MUST NOT see Record A"
    print("[SUCCESS] GET /history Data Isolation PASSED (User A & User B cannot see each other's records)")

    # 8. Test Data Isolation on GET /analysis/{id}
    # User A accesses Record A -> 200 OK
    res = client.get(f"/analysis/{record_a_id}", headers=headers_a)
    assert res.status_code == 200
    # User B attempts to access Record A -> 404 Not Found
    res_unauth = client.get(f"/analysis/{record_a_id}", headers=headers_b)
    assert res_unauth.status_code == 404
    print(f"[SUCCESS] GET /analysis/{record_a_id} Isolation PASSED (User B received 404 when accessing User A's record)")

    # 9. Test Data Isolation on DELETE /analysis/{id}
    # User B attempts to delete Record A -> 404 Not Found
    del_unauth = client.delete(f"/analysis/{record_a_id}", headers=headers_b)
    assert del_unauth.status_code == 404
    # User A deletes Record A -> 200 OK
    del_auth = client.delete(f"/analysis/{record_a_id}", headers=headers_a)
    assert del_auth.status_code == 200
    print(f"[SUCCESS] DELETE /analysis/{record_a_id} Isolation PASSED (User B prevented from deleting User A's record)")

    print("\nALL AUTHENTICATION & DATA ISOLATION TESTS PASSED SUCCESSFULLY!")


if __name__ == "__main__":
    main()
